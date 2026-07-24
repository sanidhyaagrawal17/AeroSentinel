from pptx import Presentation

prs = Presentation('AeroSentinel_Template.pptx')

# We want to replace the text in Slide 7, Shape 7 while preserving format.
slide = prs.slides[6]  # Slide 7 (0-indexed)
shape = slide.shapes[7]

# Let's verify the text first
expected_text = "Transfer learning from COCO. 100-150 epochs, batch 16, \nimg 640."
actual_text = shape.text

print(f"Original text: {actual_text}")

# We must replace text at the 'Run' level to preserve font styling
for paragraph in shape.text_frame.paragraphs:
    for run in paragraph.runs:
        if "100-150 epochs" in run.text:
            run.text = run.text.replace("100-150 epochs", "100 epochs")
        if "Transfer learning" in run.text:
            run.text = run.text.replace("Transfer learning from COCO.", "Google Colab T4 GPU. Transfer learning.")

print(f"Updated text: {shape.text}")

# Let's also check for any other '100-150' across all slides just in case
for s in prs.slides:
    for shp in s.shapes:
        if shp.has_text_frame:
            for p in shp.text_frame.paragraphs:
                for r in p.runs:
                    if "100-150" in r.text:
                        r.text = r.text.replace("100-150", "100")

output_path = 'AeroSentinel_Presentation.pptx'
prs.save(output_path)
print(f"Saved updated presentation to {output_path}")
